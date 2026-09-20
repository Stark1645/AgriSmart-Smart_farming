import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def create_element(name):
    return OxmlElement(name)

def set_cell_background(cell, fill_color):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=120, bottom=120, left=180, right=180):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_placeholder_box(doc, title, subtitle):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.rows[0].cells[0]
    set_cell_background(cell, "F8FAFC")
    set_cell_margins(cell, top=200, bottom=200, left=250, right=250)
    
    tcPr = cell._element.get_or_add_tcPr()
    tcBorders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="dashed" w:sz="12" w:space="0" w:color="94A3B8"/>'
        f'<w:left w:val="dashed" w:sz="12" w:space="0" w:color="94A3B8"/>'
        f'<w:bottom w:val="dashed" w:sz="12" w:space="0" w:color="94A3B8"/>'
        f'<w:right w:val="dashed" w:sz="12" w:space="0" w:color="94A3B8"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(tcBorders)

    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(4)
    
    r_t = p.add_run(f"[{title}]\n")
    r_t.font.bold = True
    r_t.font.size = Pt(11)
    r_t.font.color.rgb = RGBColor(15, 23, 42)
    
    r_s = p.add_run(subtitle)
    r_s.font.italic = True
    r_s.font.size = Pt(9.5)
    r_s.font.color.rgb = RGBColor(100, 116, 139)
    
    doc.add_paragraph()

def build_word_doc(output_path):
    doc = Document()
    
    # Page Margins
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Styles
    style_normal = doc.styles['Normal']
    font = style_normal.font
    font.name = 'Calibri'
    font.size = Pt(11)
    font.color.rgb = RGBColor(30, 41, 59)

    # Cover Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_badge = p_title.add_run("APPLICATION DEVELOPMENT PROJECT\n")
    r_badge.font.bold = True
    r_badge.font.size = Pt(11)
    r_badge.font.color.rgb = RGBColor(5, 150, 105)

    r_title = p_title.add_run("AgriSmart: Smart Farming & Precision Agriculture System\n")
    r_title.font.bold = True
    r_title.font.size = Pt(20)
    r_title.font.color.rgb = RGBColor(4, 120, 87)

    r_sub = p_title.add_run("Frontend & Integration Review Documentation (Review Date: Tuesday, 8th September)")
    r_sub.font.italic = True
    r_sub.font.size = Pt(12)
    r_sub.font.color.rgb = RGBColor(100, 116, 139)

    doc.add_paragraph()

    # Meta Table
    meta_table = doc.add_table(rows=3, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_table.autofit = False

    meta_data = [
        [("Student Name", "RUTHRAGURUBARAN J"), ("Register No", "727824TUIT157")],
        [("Degree & Branch", "B.TECH IT (3rd IT - 'F')"), ("Review Date", "Tuesday, 8th September")],
        [("Project Title", "AgriSmart — Smart Farming System"), ("Tech Stack", "ReactJS 19 + Spring Boot 3 + MySQL 8.x")]
    ]

    for row_idx, row in enumerate(meta_table.rows):
        for col_idx, cell in enumerate(row.cells):
            label, val = meta_data[row_idx][col_idx]
            set_cell_background(cell, "F1F5F9")
            set_cell_margins(cell, top=100, bottom=100, left=150, right=150)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r_l = p.add_run(f"{label}: ")
            r_l.font.bold = True
            r_l.font.size = Pt(10)
            r_l.font.color.rgb = RGBColor(71, 85, 105)
            r_v = p.add_run(val)
            r_v.font.size = Pt(10)
            r_v.font.color.rgb = RGBColor(15, 23, 42)

    doc.add_paragraph()

    # Section 1: Overview
    h1 = doc.add_heading("1. Executive Summary & Architecture", level=1)
    h1.runs[0].font.color.rgb = RGBColor(4, 120, 87)
    doc.add_paragraph(
        "AgriSmart is an end-to-end full-stack Precision Agriculture management platform. "
        "It features a 3-tier architecture with a ReactJS single-page frontend communicating via an "
        "interceptor-equipped Axios HTTP client to a Spring Boot 3 REST API backed by a MySQL relational database."
    )

    # Section 2: Frontend Documentation
    h2 = doc.add_heading("2. Frontend Documentation (ReactJS)", level=1)
    h2.runs[0].font.color.rgb = RGBColor(4, 120, 87)

    # 2.1 Folder Structure
    doc.add_heading("2.1 Folder Structure", level=2)
    doc.add_paragraph("Organized React directory structure separating components, pages, routing, state, services, and scoped CSS modules.")
    add_placeholder_box(doc, "SCREENSHOT 1: FOLDER STRUCTURE", "Paste screenshot of VS Code Explorer showing frontend/src, public, node_modules, etc.")

    # 2.2 Entry HTML
    doc.add_heading("2.2 Entry HTML Template (index.html)", level=2)
    doc.add_paragraph("Contains root DOM container, viewport settings, and Google Fonts preconnect links.")
    add_placeholder_box(doc, "SCREENSHOT 2: index.html", "Paste screenshot of frontend/index.html code here.")

    # 2.3 Bootstrap & Routing
    doc.add_heading("2.3 Application Bootstrap & Routing Setup", level=2)
    doc.add_paragraph("React 19 root mounting in main.jsx and master route definitions in App.jsx.")
    add_placeholder_box(doc, "SCREENSHOT 3: APPLICATION BOOTSTRAP (src/main.jsx)", "Paste screenshot of src/main.jsx code here.")
    add_placeholder_box(doc, "SCREENSHOT 4: MASTER ROUTING TABLE (src/App.jsx)", "Paste screenshot of src/App.jsx code here.")

    # 2.4 Components
    doc.add_heading("2.4 Reusable Component Architecture", level=2)
    doc.add_paragraph("Modular components including Navbar.jsx, Sidebar.jsx, StatCard.jsx, and StatusBadge.jsx.")
    add_placeholder_box(doc, "SCREENSHOT 5: NAVBAR COMPONENT (src/components/Navbar.jsx)", "Paste screenshot of Navbar.jsx code here.")
    add_placeholder_box(doc, "SCREENSHOT 6: SIDEBAR COMPONENT (src/components/Sidebar.jsx)", "Paste screenshot of Sidebar.jsx code here.")
    add_placeholder_box(doc, "SCREENSHOT 7: REUSABLE STAT CARD (src/components/StatCard.jsx)", "Paste screenshot of StatCard.jsx code here.")

    # 2.5 Pages
    doc.add_heading("2.5 Page Views & Management Modules", level=2)
    doc.add_paragraph("Application pages covering farmer command center, farm parcels, crop schedules, drone surveys, and system administration.")
    add_placeholder_box(doc, "SCREENSHOT 8: DASHBOARD PAGE (src/pages/Dashboard.jsx)", "Paste screenshot of Dashboard.jsx code here.")
    add_placeholder_box(doc, "SCREENSHOT 9: FARM MANAGEMENT PAGE (src/pages/FarmManagement.jsx)", "Paste screenshot of FarmManagement.jsx code here.")
    add_placeholder_box(doc, "SCREENSHOT 10: CROP PLANNING PAGE (src/pages/CropPlanning.jsx)", "Paste screenshot of CropPlanning.jsx code here.")
    add_placeholder_box(doc, "SCREENSHOT 11: DRONE MONITORING PAGE (src/pages/DroneMonitoring.jsx)", "Paste screenshot of DroneMonitoring.jsx code here.")
    add_placeholder_box(doc, "SCREENSHOT 12: ADMIN DASHBOARD PAGE (src/pages/AdminDashboard.jsx)", "Paste screenshot of AdminDashboard.jsx code here.")
    add_placeholder_box(doc, "SCREENSHOT 13: 3-STEP REGISTRATION PAGE (src/pages/RegisterPage.jsx)", "Paste screenshot of RegisterPage.jsx code here.")

    # 2.6 Styling
    doc.add_heading("2.6 CSS Design System & Scoped Modules", level=2)
    doc.add_paragraph("Standardized styling system using variables.css (design tokens), global.css (utilities), and PageShared.module.css (shared cards/tables).")
    add_placeholder_box(doc, "SCREENSHOT 14: CSS DESIGN TOKENS (src/styles/variables.css)", "Paste screenshot of variables.css code here.")
    add_placeholder_box(doc, "SCREENSHOT 15: GLOBAL STYLESHEET (src/styles/global.css)", "Paste screenshot of global.css code here.")
    add_placeholder_box(doc, "SCREENSHOT 16: SHARED MODULE STYLES (src/styles/PageShared.module.css)", "Paste screenshot of PageShared.module.css code here.")

    # Section 3: Integration Documentation
    h3_sec = doc.add_heading("3. Integration Documentation (React + Spring Boot + MySQL)", level=1)
    h3_sec.runs[0].font.color.rgb = RGBColor(4, 120, 87)

    doc.add_heading("3.1 Environment Configuration (.env)", level=2)
    doc.add_paragraph("Dynamic environment configuration defining VITE_API_BASE_URL=http://localhost:8080/api.")
    add_placeholder_box(doc, "SCREENSHOT 17: ENVIRONMENT CONFIG (frontend/.env)", "Paste screenshot of frontend/.env code here.")

    doc.add_heading("3.2 Axios HTTP Client & Interceptors (services/api.js)", level=2)
    doc.add_paragraph("Axios client setup with JWT Bearer token request interceptor, unified response error handler, and REST service methods.")
    add_placeholder_box(doc, "SCREENSHOT 18: AXIOS INSTANCE & INTERCEPTORS (src/services/api.js)", "Paste screenshot of Axios interceptor code in api.js here.")
    add_placeholder_box(doc, "SCREENSHOT 19: AXIOS REST API SERVICES (src/services/api.js)", "Paste screenshot of service functions (authAPI, farmAPI, sensorAPI, cropAPI, etc.) in api.js here.")

    doc.add_heading("3.3 Spring Boot REST Controllers & MySQL Configuration", level=2)
    doc.add_paragraph("Backend Java controllers and database connection configuration.")
    add_placeholder_box(doc, "SCREENSHOT 20: AUTH CONTROLLER (backend/.../controller/AuthController.java)", "Paste screenshot of AuthController.java code here.")
    add_placeholder_box(doc, "SCREENSHOT 21: FARM CONTROLLER (backend/.../controller/FarmController.java)", "Paste screenshot of FarmController.java code here.")
    add_placeholder_box(doc, "SCREENSHOT 22: MYSQL DATABASE PROPERTIES (backend/.../application.properties)", "Paste screenshot of application.properties code here.")

    # Section 4: Live UI
    h4_sec = doc.add_heading("4. Live Running Application Interface", level=1)
    h4_sec.runs[0].font.color.rgb = RGBColor(4, 120, 87)
    doc.add_paragraph("Live dashboard interface running on http://localhost:5173 connected to Spring Boot backend on port 8080.")
    add_placeholder_box(doc, "SCREENSHOT 23: LIVE APPLICATION DASHBOARD IN BROWSER", "Paste screenshot of live browser UI at http://localhost:5173/dashboard here.")

    # Save
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    doc.save(output_path)
    print(f"Successfully updated: {output_path}")

if __name__ == '__main__':
    output_docx = r"e:\New folder\AgriSmart-Smart_farming-master\docs\AgriSmart_Frontend_and_Integration_Documentation.docx"
    build_word_doc(output_docx)
